"""
JobNex Resume Builder - Phase 8 Automated Test Suite
Verifies:
1. ATS Word Document (.docx) export endpoint and document formatting via python-docx
2. Clean ATS Plain-Text (.txt) export endpoint and uppercase section structure
3. Custom filename sanitization and query parameter support
4. Live/draft POST export support with client-side overrides
5. Ownership security (403 for unauthorized users, 404 for nonexistent resumes)
6. Zero placeholder verification in generated document text
"""
import io
import pytest
import docx
from fastapi.testclient import TestClient

from main import app
from resume_builder import (
    ResumeData,
    PersonalInfo,
    PersonalLink,
    ResumeSection,
    ResumeItem,
    get_export_filename,
    generate_ats_word_document,
    generate_clean_plain_text,
)


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def auth_header_user1():
    return {"X-User-Id": "u-test-phase8"}


@pytest.fixture
def auth_header_user2():
    return {"X-User-Id": "u-other-phase8"}


@pytest.fixture
def sample_resume(client, auth_header_user1):
    # Create sample resume for testing
    create_res = client.post(
        "/api/resumes",
        json={"name": "Alex Mercer - Senior Cloud Architect"},
        headers=auth_header_user1,
    )
    assert create_res.status_code == 201
    res_id = create_res.json()["id"]

    # Populate resume with detailed ATS content
    update_payload = {
        "content": {
            "name": "Alex Mercer - Senior Cloud Architect",
            "personal": {
                "name": "Alex Mercer",
                "email": "alex.mercer@example.com",
                "phone": "+1 (555) 019-9832",
                "location": "Seattle, WA",
                "links": [
                    {"label": "LinkedIn", "url": "https://linkedin.com/in/alex-mercer"},
                    {"label": "GitHub", "url": "https://github.com/alex-mercer"},
                ],
            },
            "summary": "Results-driven Cloud Solutions Architect with 8+ years designing scalable microservices.",
            "sections": [
                {
                    "id": "sec-work-1",
                    "type": "work",
                    "title": "Work Experience",
                    "hidden": False,
                    "items": [
                        {
                            "id": "item-w1",
                            "title": "Lead Cloud Architect",
                            "organization": "Nexus Cloud Systems",
                            "dates": "2021 - Present",
                            "location": "Seattle, WA",
                            "bullets": [
                                "Architected AWS multi-region infrastructure handling 50M daily transactions.",
                                "Automated Kubernetes CI/CD deployment pipelines cutting lead time by 45%.",
                            ],
                            "technologies": ["AWS", "Kubernetes", "Terraform", "Python"],
                            "hidden": False,
                        }
                    ],
                },
                {
                    "id": "sec-edu-1",
                    "type": "education",
                    "title": "Education",
                    "hidden": False,
                    "items": [
                        {
                            "id": "item-e1",
                            "title": "B.S. in Computer Science",
                            "organization": "University of Washington",
                            "dates": "2013 - 2017",
                            "location": "Seattle, WA",
                            "bullets": [],
                            "technologies": [],
                            "hidden": False,
                        }
                    ],
                },
                {
                    "id": "sec-skills-1",
                    "type": "skills",
                    "title": "Technical Skills",
                    "hidden": False,
                    "items": [
                        {
                            "id": "item-sk1",
                            "title": "Cloud & DevOps",
                            "organization": "",
                            "dates": "",
                            "location": "",
                            "bullets": [],
                            "technologies": ["AWS", "Docker", "Kubernetes", "Terraform"],
                            "hidden": False,
                        },
                        {
                            "id": "item-sk2",
                            "title": "Languages & Frameworks",
                            "organization": "",
                            "dates": "",
                            "location": "",
                            "bullets": [],
                            "technologies": ["Python", "FastAPI", "Go", "TypeScript"],
                            "hidden": False,
                        },
                    ],
                },
            ],
        }
    }
    up_res = client.put(f"/api/resumes/{res_id}", json=update_payload, headers=auth_header_user1)
    assert up_res.status_code == 200
    return res_id


def test_filename_sanitizer():
    dummy = ResumeData(
        personal=PersonalInfo(name="Alex O'Connor / Architect"),
        name="Alex Resume",
    )
    fname_docx = get_export_filename(dummy, None, ".docx")
    assert fname_docx == "Alex_OConnor_Architect_Resume.docx"

    fname_custom = get_export_filename(dummy, "My Custom Resume 2026!", ".pdf")
    assert fname_custom == "My_Custom_Resume_2026.pdf"

    # Fallback when empty
    empty_resume = ResumeData(personal=PersonalInfo(name=""), name="Untitled Resume")
    fname_fallback = get_export_filename(empty_resume, None, ".docx")
    assert fname_fallback == "JobNex_Resume.docx"


def test_export_docx_endpoint_format_and_structure(client, auth_header_user1, sample_resume):
    response = client.get(f"/api/resumes/{sample_resume}/export-docx", headers=auth_header_user1)
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers["Content-Type"]
    assert "attachment" in response.headers["Content-Disposition"]
    assert "Alex_Mercer_Resume.docx" in response.headers["Content-Disposition"]

    # Verify binary zip header of DOCX format
    raw_bytes = response.content
    assert raw_bytes.startswith(b"PK\x03\x04")

    # Parse and inspect with python-docx
    doc = docx.Document(io.BytesIO(raw_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])

    assert "ALEX MERCER" in full_text.upper()
    assert "alex.mercer@example.com" in full_text
    assert "Seattle, WA" in full_text
    assert "PROFESSIONAL SUMMARY" in full_text
    assert "Results-driven Cloud Solutions Architect" in full_text
    assert "WORK EXPERIENCE" in full_text
    assert "Lead Cloud Architect" in full_text
    assert "Nexus Cloud Systems" in full_text
    assert "50M daily transactions" in full_text
    assert "EDUCATION" in full_text
    assert "University of Washington" in full_text
    assert "TECHNICAL SKILLS" in full_text
    assert "Cloud & DevOps" in full_text
    assert "AWS, Docker, Kubernetes, Terraform" in full_text

    # Verify no bracketed placeholders
    assert "[" not in full_text
    assert "]" not in full_text


def test_export_txt_endpoint_format_and_structure(client, auth_header_user1, sample_resume):
    response = client.get(f"/api/resumes/{sample_resume}/export-txt", headers=auth_header_user1)
    assert response.status_code == 200
    assert "text/plain" in response.headers["Content-Type"]
    assert "attachment" in response.headers["Content-Disposition"]
    assert "Alex_Mercer_Resume.txt" in response.headers["Content-Disposition"]

    text = response.text
    assert "ALEX MERCER" in text
    assert "alex.mercer@example.com" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert "WORK EXPERIENCE" in text
    assert "EDUCATION" in text
    assert "TECHNICAL SKILLS" in text
    assert "Lead Cloud Architect" in text
    assert "50M daily transactions" in text

    # Verify no bracket placeholders
    assert "[" not in text
    assert "]" not in text


def test_export_custom_filename_query_param(client, auth_header_user1, sample_resume):
    res_docx = client.get(
        f"/api/resumes/{sample_resume}/export-docx?filename=Alex_Mercer_DevOps_2026",
        headers=auth_header_user1,
    )
    assert res_docx.status_code == 200
    assert 'filename="Alex_Mercer_DevOps_2026.docx"' in res_docx.headers["Content-Disposition"]

    res_txt = client.get(
        f"/api/resumes/{sample_resume}/export-txt?filename=Alex_Mercer_DevOps_2026",
        headers=auth_header_user1,
    )
    assert res_txt.status_code == 200
    assert 'filename="Alex_Mercer_DevOps_2026.txt"' in res_txt.headers["Content-Disposition"]


def test_export_post_live_draft_override(client, auth_header_user1, sample_resume):
    # Overwrite summary in live draft without saving to DB first
    draft_payload = {
        "personal": {"name": "Alex Mercer Live Draft", "email": "live@example.com"},
        "summary": "Unsaved live draft summary for testing immediate export fidelity.",
        "filename": "Alex_Mercer_Live_Export",
    }
    res_post = client.post(
        f"/api/resumes/{sample_resume}/export-docx",
        json=draft_payload,
        headers=auth_header_user1,
    )
    assert res_post.status_code == 200
    assert 'filename="Alex_Mercer_Live_Export.docx"' in res_post.headers["Content-Disposition"]

    doc = docx.Document(io.BytesIO(res_post.content))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Alex Mercer Live Draft" in full_text
    assert "Unsaved live draft summary" in full_text


def test_export_security_and_nonexistent(client, auth_header_user1, auth_header_user2, sample_resume):
    # 404 for non-existent resume
    res_404 = client.get("/api/resumes/res-nonexistent-999/export-docx", headers=auth_header_user1)
    assert res_404.status_code == 404

    # 403 for unauthorized other user
    res_403 = client.get(f"/api/resumes/{sample_resume}/export-docx", headers=auth_header_user2)
    assert res_403.status_code == 403
